import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from src.lab04.io_txt_csv import ensure_parent_dir


def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    """Конвертирует CSV в XLSX.

    Первая строка CSV — заголовок. Лист называется "Sheet1". 
    Ширина каждой колонки равна длине самого длинного значения, но не меньше 8 символов.
    Значения записываются как строки (типы не угадываются).

    Raises:
        FileNotFoundError: входного CSV нет.
        ValueError: неверное расширение, пустой CSV.
    """
    src = Path(csv_path)
    dst = Path(xlsx_path)
    if src.suffix.lower() != ".csv":
        raise ValueError(f"Неверный тип файла: ожидается .csv, получено '{src.suffix}'")
    if dst.suffix.lower() != ".xlsx":
        raise ValueError(f"Неверный тип файла: ожидается .xlsx, получено '{dst.suffix}'")
    if not src.is_file():
        raise FileNotFoundError(f"Файл не найден: {csv_path}")

    rows = []
    with open(csv_path, encoding="utf-8", newline="") as f:
        for row in csv.reader(f):
            if row:
                rows.append(row)

    # нужна хотя бы строка заголовка и одна строка данных
    if len(rows) < 2:
        raise ValueError("CSV пустой или в нём нет данных (только заголовок)")

    widths = {}
    for row in rows:
        for index, cell in enumerate(row):
            if len(cell) > widths.get(index, 0):
                widths[index] = len(cell)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for row in rows:
        ws.append(row)

    for index, width in widths.items():
        letter = get_column_letter(index + 1)
        ws.column_dimensions[letter].width = max(8, width)

    ensure_parent_dir(dst)
    wb.save(dst)

# csv_to_xlsx("data/samples/people.csv", "data/out/people.xlsx")
# csv_to_xlsx("data/samples/cities.csv", "data/out/cities.xlsx")